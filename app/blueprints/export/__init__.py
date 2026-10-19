"""엑셀 내려받기 (admin). 현재 필터 그대로, access_logs 기록 (CLAUDE.md 7-7, 9)."""
from io import BytesIO

from flask import Blueprint, request, send_file
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ... import now_kst
from ...extensions import db
from ...models import STATUS_LABELS, AccessLog, ConsultSheet, Inquiry
from ...permissions import role_required
from ...services import inquiry as svc
from ...services import sheet as sheet_svc

bp = Blueprint("export", __name__)

HEAD = ["번호", "상태", "접수일시", "사이트", "업종", "고객명", "연락처", "지역", "문의유형", "문의내용", "담당자", "콜백",
        "유입 키워드", "utm_source", "utm_medium", "utm_campaign", "기기", "중복", "첫연락", "계약일", "종결일", "종결사유"]


@bp.get("/inquiries/export.xlsx")
@role_required("admin")
def export_xlsx():
    base, q = svc.apply_filters(svc.scoped_query(current_user), current_user, request.args)
    rows = q.order_by(Inquiry.received_at.desc()).all()
    sheets = {s.inquiry_id: s for s in ConsultSheet.query.filter(ConsultSheet.is_current.is_(True),
                                                                   ConsultSheet.inquiry_id.in_([r.id for r in rows] or [-1]))}
    form_ids = {s.form_id for s in sheets.values()}
    single_form = sheets[next(iter(sheets))].form if len(form_ids) == 1 else None  # 양식이 하나로 통일될 때만 열 펼침

    wb = Workbook()
    ws = wb.active
    ws.title = "문의"
    head = list(HEAD)
    if single_form:
        head += [f["label"] for f in sheet_svc.fields_of(single_form.schema_json)]
    elif sheets:
        head.append("상담지 요약")
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF0F3")

    f = lambda d: d.strftime("%Y-%m-%d %H:%M") if d else ""
    for r in rows:
        row = [r.id, STATUS_LABELS[r.status], f(r.received_at), r.site.name, {"legal": "법률", "labor": "인사·노무"}[r.site.category],
               r.customer_name, r.phone, r.region or "", r.inquiry_type or "", r.message or "",
               r.assignee.display if r.assignee else "미배정", f(r.callback_at), r.utm_term or "", r.utm_source or "",
               r.utm_medium or "", r.utm_campaign or "", r.device or "", "Y" if r.duplicate_of_id else "",
               f(r.first_contact_at), f(r.contracted_at), f(r.closed_at), r.close_reason or ""]
        s = sheets.get(r.id)
        if single_form:
            a = s.answers_json if s else {}
            for fld in sheet_svc.fields_of(single_form.schema_json):
                v = a.get(fld["key"], "")
                row.append(", ".join(v) if isinstance(v, list) else ("예" if v is True else "아니오" if v is False else v))
        elif sheets:
            row.append(sheet_svc.summary_text(s) if s else "")
        ws.append(row)
    for i, _ in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.column_dimensions["J"].width = 50

    db.session.add(AccessLog(user_id=current_user.id, action="export_excel", ip=request.remote_addr, created_at=now_kst(),
                             detail=f"{len(rows)}건 · " + "&".join(f"{k}={v}" for k, v in request.args.items() if v)[:200]))
    db.session.commit()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"문의_{now_kst():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
